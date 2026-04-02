from typing import List, Dict
from docx import Document

class AssemblerAgent:
    def __init__(self, mock_mode: bool = False):
        self.mock_mode = mock_mode

    def export_to_docx(self, chapters: List[Dict[str, str]], filename: str):
        if self.mock_mode:
            print(f"Mock export: Would save {len(chapters)} chapters to {filename}")
            return
        
        doc = Document()
        doc.add_heading("Course Syllabus", level=1)
        for chapter in chapters:
            doc.add_heading(chapter["title"], level=2)
            doc.add_paragraph(chapter["content"])
        doc.save(filename)